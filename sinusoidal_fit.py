import sys
import re
from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from scipy.interpolate import CubicSpline
from scipy.optimize import curve_fit
from scipy.signal import butter, find_peaks, iirnotch, sosfiltfilt, filtfilt, sosfilt, resample_poly
from scipy.ndimage import median_filter
from scipy.stats import median_abs_deviation

def sinusoid(x, A, B, C, D):
    """Sinusoidal model: y = A * sin(B*x + C) + D."""
    return A * np.sin(B * x + C) + D


def estimate_initial_guess(x, y):
    """Estimate starting values for nonlinear least squares fitting."""
    if len(x) < 4:
        raise ValueError("At least 4 numeric x,y data points are required.")

    amplitude_guess = (np.max(y) - np.min(y)) / 2
    offset_guess = np.mean(y)

    # Estimate angular frequency B using the FFT.
    x_differences = np.diff(x)
    nonzero_differences = np.abs(x_differences[x_differences != 0])

    if len(nonzero_differences) == 0:
        raise ValueError("The x values must not all be the same.")

    x_spacing = np.mean(nonzero_differences)
    y_centered = y - offset_guess

    frequencies = np.fft.rfftfreq(len(x), d=x_spacing)
    fft_magnitudes = np.abs(np.fft.rfft(y_centered))

    # Ignore the zero-frequency term when finding the dominant frequency.
    if len(fft_magnitudes) > 1 and np.any(fft_magnitudes[1:] > 0):
        dominant_index = np.argmax(fft_magnitudes[1:]) + 1
        frequency_cycles = frequencies[dominant_index]
        angular_frequency_guess = 2 * np.pi * frequency_cycles
    else:
        angular_frequency_guess = 2 * np.pi / (np.max(x) - np.min(x))

    phase_guess = 0

    return amplitude_guess, angular_frequency_guess, phase_guess, offset_guess


def fit_sinusoidal(x_values, y_values):
    """Fit x,y data to y = A * sin(B*x + C) + D."""
    x_values, y_values = prepare_xy_values(x_values, y_values)

    initial_guess = estimate_initial_guess(x_values, y_values)

    fitted_parameters, covariance = curve_fit(
        sinusoid,
        x_values,
        y_values,
        p0=initial_guess,
        maxfev=10000,
    )

    A, B, C, D = fitted_parameters

    # Make the printed equation easier to read by keeping amplitude positive.
    if A < 0:
        A = -A
        C += np.pi

    # Normalize phase into the range [-pi, pi].
    C = (C + np.pi) % (2 * np.pi) - np.pi

    return A, B, C, D


def prepare_xy_values(x_values, y_values):
    """Sort x,y values and ensure each x value has one y value."""
    x_values = np.asarray(x_values, dtype=float)
    y_values = np.asarray(y_values, dtype=float)

    sorted_indices = np.argsort(x_values)
    x_values = x_values[sorted_indices]
    y_values = y_values[sorted_indices]

    unique_x = []
    unique_y = []

    for x_value, y_value in zip(x_values, y_values):
        if unique_x and np.isclose(x_value, unique_x[-1]):
            if not np.isclose(y_value, unique_y[-1]):
                raise ValueError(
                    "An exact y=f(x) formula is impossible because the same "
                    f"x value ({x_value}) has more than one y value."
                )
            continue

        unique_x.append(x_value)
        unique_y.append(y_value)

    if len(unique_x) < 4:
        raise ValueError("At least 4 unique numeric x,y data points are required.")

    return np.array(unique_x), np.array(unique_y)


def make_exact_spline(x_values, y_values):
    """Create a cubic spline that passes exactly through every data point."""
    x_values, y_values = prepare_xy_values(x_values, y_values)
    spline = CubicSpline(x_values, y_values, bc_type="natural")
    return x_values, y_values, spline


def calculate_spline_area(spline, x_min, x_max):
    """Calculate signed and absolute area under an exact cubic spline."""
    signed_area = float(spline.integrate(x_min, x_max))

    roots = spline.roots(extrapolate=False)
    roots = roots[np.isreal(roots)].real
    roots = roots[(roots > x_min) & (roots < x_max)]

    split_points = np.concatenate(([x_min], np.sort(roots), [x_max]))
    absolute_area = 0

    for start, end in zip(split_points[:-1], split_points[1:]):
        absolute_area += abs(float(spline.integrate(start, end)))

    return signed_area, absolute_area


def calculate_sinusoid_area(A, B, C, D, x_min, x_max):
    """Calculate the signed area under y = A*sin(B*x + C) + D."""
    if np.isclose(B, 0):
        return float((A * np.sin(C) + D) * (x_max - x_min))

    start_value = (-A / B) * np.cos(B * x_min + C) + D * x_min
    end_value = (-A / B) * np.cos(B * x_max + C) + D * x_max

    return float(end_value - start_value)


def calculate_sinusoid_absolute_area(A, B, C, D, x_min, x_max):
    """Calculate absolute area under y = A*sin(B*x + C) + D."""
    if x_max < x_min:
        x_min, x_max = x_max, x_min

    if np.isclose(A, 0) or np.isclose(B, 0):
        return abs(calculate_sinusoid_area(A, B, C, D, x_min, x_max))

    target = -D / A
    roots = []

    if abs(target) <= 1:
        base_angle = np.arcsin(target)
        candidate_angles = [base_angle, np.pi - base_angle]
        theta_start = B * x_min + C
        theta_end = B * x_max + C
        theta_min = min(theta_start, theta_end)
        theta_max = max(theta_start, theta_end)
        k_min = int(np.floor((theta_min - 2 * np.pi) / (2 * np.pi))) - 1
        k_max = int(np.ceil((theta_max + 2 * np.pi) / (2 * np.pi))) + 1

        for k in range(k_min, k_max + 1):
            for angle in candidate_angles:
                root = (angle + 2 * np.pi * k - C) / B
                if x_min < root < x_max:
                    roots.append(root)

    split_points = np.array([x_min, *sorted(set(np.round(roots, 12))), x_max])
    absolute_area = 0

    for start, end in zip(split_points[:-1], split_points[1:]):
        absolute_area += abs(calculate_sinusoid_area(A, B, C, D, start, end))

    return float(absolute_area)


def _interpolated_x_at_level(x0, y0, x1, y1, level):
    """Return linearly interpolated x where y crosses level."""
    if np.isclose(y1, y0):
        return float(x0)

    fraction = (level - y0) / (y1 - y0)
    return float(x0 + fraction * (x1 - x0))


def _find_left_crossing(x_values, signal, peak_index, level):
    for index in range(peak_index - 1, -1, -1):
        if signal[index] <= level <= signal[index + 1]:
            return _interpolated_x_at_level(
                x_values[index],
                signal[index],
                x_values[index + 1],
                signal[index + 1],
                level,
            )

    return None


def _find_right_crossing(x_values, signal, peak_index, level):
    for index in range(peak_index, len(signal) - 1):
        if signal[index] >= level >= signal[index + 1]:
            return _interpolated_x_at_level(
                x_values[index],
                signal[index],
                x_values[index + 1],
                signal[index + 1],
                level,
            )

    return None


def calculate_spike_metrics(x_values, y_values):
    """Calculate dominant epileptiform spike morphology from x,y trace data."""
    x_values, y_values = prepare_xy_values(x_values, y_values)

    baseline = float(np.median(y_values))
    positive_peak_index = int(np.argmax(y_values - baseline))
    negative_peak_index = int(np.argmax(baseline - y_values))

    positive_amplitude = float(y_values[positive_peak_index] - baseline)
    negative_amplitude = float(baseline - y_values[negative_peak_index])

    if positive_amplitude >= negative_amplitude:
        peak_index = positive_peak_index
        polarity = "positive"
        signed_amplitude = positive_amplitude
        signal = y_values - baseline
    else:
        peak_index = negative_peak_index
        polarity = "negative"
        signed_amplitude = -negative_amplitude
        signal = baseline - y_values

    peak_amplitude = float(abs(signed_amplitude))
    peak_x = float(x_values[peak_index])
    peak_y = float(y_values[peak_index])

    if np.isclose(peak_amplitude, 0):
        raise ValueError("Could not calculate spike metrics because peak amplitude is 0.")

    level_10 = peak_amplitude * 0.10
    level_90 = peak_amplitude * 0.90

    rise_10_x = _find_left_crossing(x_values, signal, peak_index, level_10)
    rise_90_x = _find_left_crossing(x_values, signal, peak_index, level_90)
    decay_90_x = _find_right_crossing(x_values, signal, peak_index, level_90)
    decay_10_x = _find_right_crossing(x_values, signal, peak_index, level_10)

    rise_time = None
    if rise_10_x is not None and rise_90_x is not None:
        rise_time = float(rise_90_x - rise_10_x)

    decay_time = None
    if decay_90_x is not None and decay_10_x is not None:
        decay_time = float(decay_10_x - decay_90_x)

    duration_10_to_10 = None
    if rise_10_x is not None and decay_10_x is not None:
        duration_10_to_10 = float(decay_10_x - rise_10_x)

    return {
        "Baseline y (median)": baseline,
        "Spike polarity": polarity,
        "Peak x": peak_x,
        "Peak y": peak_y,
        "Peak amplitude (signed from baseline)": signed_amplitude,
        "Peak amplitude (absolute from baseline)": peak_amplitude,
        "Rise start x (10% amplitude)": rise_10_x,
        "Rise end x (90% amplitude)": rise_90_x,
        "Rise time (10-90%)": rise_time,
        "Decay start x (90% amplitude)": decay_90_x,
        "Decay end x (10% amplitude)": decay_10_x,
        "Decay time (90-10%)": decay_time,
        "Spike duration (10-10%)": duration_10_to_10,
        "Measurement note": (
            "Baseline is median y. Peak is largest absolute deflection from baseline. "
            "Rise and decay use linear interpolation between sampled x,y points."
        ),
    }


def build_spike_metric_rows(spike_metrics):
    """Build rows for the Spike Metrics export sheet."""
    rows = []

    for metric, value in spike_metrics.items():
        rows.append(
            {
                "Metric": metric,
                "Value": "" if value is None else value,
            }
        )

    return rows


def infer_sampling_rate_from_time(x_values):
    """Infer sampling rate from x values, converting milliseconds to seconds if needed."""
    x_values = np.asarray(x_values, dtype=float)
    if len(x_values) < 3:
        raise ValueError("At least 3 time points are required.")

    dt = np.diff(x_values)
    dt = dt[np.isfinite(dt) & (dt > 0)]
    if len(dt) == 0:
        raise ValueError("Could not infer sampling rate from time column.")

    dt_median = float(np.median(dt))
    time_scale = 1.0
    fs = 1.0 / dt_median

    if fs < 50:
        time_scale = 0.001
        fs = 1.0 / (dt_median * time_scale)

    return fs, x_values * time_scale


def load_tsv_waveform(file_path):
    """Load TSV/TXT/CSV waveform data, preferring columns D(time) and E(signal)."""
    path = Path(file_path)
    delimiter = "," if path.suffix.lower() == ".csv" else "\t"

    rows = []
    with open(path, "r", encoding="utf-8", errors="ignore") as file:
        for line in file:
            parts = line.rstrip("\n").split(delimiter)
            if len(parts) >= 5:
                rows.append((parts[3], parts[4]))
            elif len(parts) >= 2:
                rows.append((parts[0], parts[1]))

    points = []
    skipped = 0
    for x_raw, y_raw in rows:
        x_value = to_float(x_raw)
        y_value = to_float(y_raw)
        if x_value is None or y_value is None:
            skipped += 1
            continue
        points.append((x_value, y_value))

    if len(points) < 20:
        raise ValueError(f"{path.name}: not enough numeric waveform points.")

    points.sort(key=lambda point: point[0])
    x_values = np.array([point[0] for point in points], dtype=float)
    y_values = np.array([point[1] for point in points], dtype=float)
    fs, t_seconds = infer_sampling_rate_from_time(x_values)

    return t_seconds, y_values, fs, skipped


def iter_tsv_waveform_chunks(file_path, chunk_rows=600000):
    """Yield numeric time/signal chunks from TSV/TXT/CSV files."""
    path = Path(file_path)
    delimiter = "," if path.suffix.lower() == ".csv" else "\t"
    chunk_rows = int(chunk_rows)
    x_values = []
    y_values = []
    skipped = 0
    previous_fs = None
    previous_time_scale = None

    def build_chunk():
        nonlocal previous_fs, previous_time_scale
        x_chunk = np.asarray(x_values, dtype=float)
        y_chunk = np.asarray(y_values, dtype=float)

        if len(x_chunk) >= 3:
            fs, t_seconds = infer_sampling_rate_from_time(x_chunk)
            previous_fs = fs
            if np.any(x_chunk != 0):
                nonzero_index = int(np.flatnonzero(x_chunk != 0)[0])
                previous_time_scale = float(t_seconds[nonzero_index] / x_chunk[nonzero_index])
            elif previous_time_scale is None:
                previous_time_scale = 1.0
            return t_seconds, y_chunk, fs

        if previous_fs is None or previous_time_scale is None:
            raise ValueError("At least 3 time points are required.")

        return x_chunk * previous_time_scale, y_chunk, previous_fs

    with open(path, "r", encoding="utf-8", errors="ignore") as file:
        for line in file:
            parts = line.rstrip("\r\n").split(delimiter)
            if len(parts) < 5:
                skipped += 1
                continue

            x_value = to_float(parts[3])
            y_value = to_float(parts[4])
            if x_value is None or y_value is None:
                skipped += 1
                continue

            x_values.append(x_value)
            y_values.append(y_value)

            if len(x_values) >= chunk_rows:
                t_seconds, y_chunk, fs = build_chunk()
                yield t_seconds, y_chunk, fs, skipped
                x_values = []
                y_values = []
                skipped = 0

    if x_values:
        t_seconds, y_chunk, fs = build_chunk()
        yield t_seconds, y_chunk, fs, skipped


def _butter_bandpass_sos(fs, low, high, order=4):
    nyquist = 0.5 * fs
    low = max(1e-6, low / nyquist)
    high = min(0.999999, high / nyquist)
    return butter(order, [low, high], btype="band", output="sos")


def _apply_notch(signal, fs, line_hz=50.0, q=35.0):
    w0 = line_hz / (fs / 2.0)
    if not (0 < w0 < 1):
        return signal

    b, a = iirnotch(w0, q)
    padlen = min(len(signal) - 1, 3 * max(len(a), len(b)))
    return filtfilt(b, a, signal, padtype="odd", padlen=padlen)


def _moving_mad_stats(values, fs, win_s=10.0):
    med_value = float(np.median(values))
    dev_value = float(median_abs_deviation(values, scale="normal")) + 1e-9
    return np.full_like(values, med_value), np.full_like(values, dev_value)


def _band_power_fft(values, fs, bands):
    n = len(values)
    freqs = np.fft.rfftfreq(n, d=1.0 / fs)
    power = (np.abs(np.fft.rfft(values)) ** 2) / max(1, n)
    return [
        float(np.sum(power[(freqs >= low) & (freqs < high)]))
        for low, high in bands
    ]


def _eeg_preprocess(signal, fs):
    width = max(3, int(round(fs * 2.0)))
    if width % 2 == 0:
        width += 1
    baseline = median_filter(signal, size=width, mode="nearest")
    processed = signal - baseline
    processed = _apply_notch(processed, fs, 50.0, q=30.0)
    processed = _apply_notch(processed, fs, 100.0, q=30.0)
    return sosfiltfilt(
        _butter_bandpass_sos(fs, 0.5, min(100.0, 0.45 * fs), order=4),
        processed,
    )


def _decimate_to(signal, fs, target_fs):
    if target_fs >= fs:
        return signal, fs

    cutoff = 0.45 * target_fs
    filtered = sosfilt(butter(6, cutoff, btype="low", fs=fs, output="sos"), signal)
    down = max(1, int(round(fs / target_fs)))
    return resample_poly(filtered, 1, down), fs / down


def detect_epileptiform_spikes(signal, fs):
    """Detect sharp epileptiform spikes using EEG/ECG analyser spike criteria."""
    signal = np.asarray(signal, dtype=float)
    signal = _eeg_preprocess(signal, fs)
    signal, fs = _decimate_to(signal, fs, 500)

    sos = _butter_bandpass_sos(fs, 14.0, min(70.0, 0.45 * fs), order=4)
    hf = sosfiltfilt(sos, signal)
    hf_abs = np.abs(hf)
    med, dev = _moving_mad_stats(hf_abs, fs, 10.0)
    z_values = (hf_abs - med) / dev

    min_dist = max(1, int(fs * 50 / 1000.0))
    peaks, _ = find_peaks(z_values, height=4.8, distance=min_dist)

    grad_med = np.median(np.abs(np.diff(signal))) + 1e-9
    curv_med = np.median(np.abs(np.diff(signal, n=2))) + 1e-9
    events = []
    n = len(signal)

    for peak in peaks:
        half = 0.5 * hf_abs[peak]
        left = right = int(peak)
        while left > 0 and hf_abs[left] > half:
            left -= 1
        while right < n - 1 and hf_abs[right] > half:
            right += 1

        width_ms = 1000.0 * max(1, right - left) / fs
        if not (8 <= width_ms <= 100):
            continue

        k = max(2, int(round(0.005 * fs)))
        core_left = max(0, peak - k)
        core_right = min(n, peak + k)
        core = signal[core_left:core_right]
        peak_slope = float(np.max(np.abs(np.diff(core)))) if core.size >= 2 else 0.0

        pre_span = max(k * 5, int(0.02 * fs))
        base_left = max(0, core_left - pre_span)
        base = signal[base_left:core_left]
        base_grad = float(np.median(np.abs(np.diff(base)))) + 1e-9 if base.size >= 2 else grad_med
        if peak_slope < 1.4 * base_grad:
            continue

        sharp = float(np.max(np.abs(np.diff(core, n=2)))) if core.size >= 3 else 0.0
        if sharp < 1.2 * curv_med:
            continue

        seg_start = max(0, peak - int(0.04 * fs))
        seg_end = min(n, peak + int(0.04 * fs))
        segment = signal[seg_start:seg_end]
        if segment.size >= int(0.02 * fs):
            high_power, low_power = _band_power_fft(segment, fs, [(30, 80), (1, 20)])
            if low_power <= 0 or (high_power / low_power) < 0.05:
                continue

        amp_window = int(round(3.0 * fs))
        amp_start = max(0, peak - amp_window // 2)
        amp_end = min(n, peak + amp_window // 2)
        seg_abs = np.abs(signal[amp_start:amp_end])
        med_abs = float(np.median(seg_abs))
        mad_abs = float(median_abs_deviation(seg_abs, scale="normal")) + 1e-9
        z_amp = (abs(float(signal[peak])) - med_abs) / mad_abs
        if z_amp < 7.0:
            continue

        events.append(
            {
                "idx": int(peak),
                "time_s": float(peak / fs),
                "width_ms": float(width_ms),
                "z_amp": float(z_amp),
                "confidence": float(1.0 / (1.0 + np.exp(-(z_amp - 7.0)))),
            }
        )

    return events


def _find_alignment_extremum(signal, event_index, fs, search_ms=60.0):
    """Find local dominant peak/trough near detector event."""
    radius = max(1, int(round(search_ms / 1000.0 * fs)))
    left = max(0, int(event_index) - radius)
    right = min(len(signal), int(event_index) + radius + 1)
    window = signal[left:right]

    if len(window) == 0:
        return int(event_index), 0.0, "unknown"

    baseline = float(np.median(window))
    relative = window - baseline
    local_index = int(np.argmax(np.abs(relative)))
    align_index = left + local_index
    amplitude = float(relative[local_index])
    polarity = "positive peak" if amplitude >= 0 else "negative trough"
    return align_index, amplitude, polarity


def _read_declared_sample_count(file_path):
    """Read Axion-style Count header from large text exports."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            for _ in range(20):
                line = file.readline()
                if not line:
                    break
                parts = [part.strip() for part in line.rstrip("\n").split("\t")]
                if not parts:
                    continue
                key = parts[0].strip().rstrip(":").lower()
                if key != "count" or len(parts) < 2:
                    continue
                return int(float(parts[1]))
    except Exception:
        return None
    return None


def _progress_percent_from_eeg_message(message, total_samples):
    if not total_samples:
        return None

    match = re.search(r"scanned\s+([\d,]+)", str(message))
    if not match:
        return None

    try:
        processed = int(match.group(1).replace(",", ""))
    except ValueError:
        return None

    return 5.0 + 85.0 * min(1.0, processed / max(1, int(total_samples)))


def _read_spike_window(eeg_module, path, start_s, end_s, fs_hint):
    if not hasattr(eeg_module, "read_waveform_window"):
        raise RuntimeError("EEG/ECG analyser read_waveform_window() is unavailable.")

    time_s, signal = eeg_module.read_waveform_window(
        str(path),
        float(start_s),
        float(end_s),
        fs_hint=fs_hint,
    )
    time_s = np.asarray(time_s, dtype=float)
    signal = np.asarray(signal, dtype=float)
    valid = np.isfinite(time_s) & np.isfinite(signal)
    if not np.any(valid):
        return None, None, None

    time_s = time_s[valid]
    signal = signal[valid]
    order = np.argsort(time_s)
    time_s = time_s[order]
    signal = signal[order]
    unique_time, unique_index = np.unique(time_s, return_index=True)
    signal = signal[unique_index]
    time_s = unique_time
    if len(time_s) < 5:
        return None, None, None

    if fs_hint is not None and np.isfinite(float(fs_hint)) and float(fs_hint) > 0:
        fs = float(fs_hint)
    else:
        fs, time_s = infer_sampling_rate_from_time(time_s)

    return time_s, signal, fs


def _extract_aligned_spike_segment(
    eeg_module,
    path,
    spike_time,
    pre_ms,
    post_ms,
    target_x_ms=None,
    fs_hint=None,
    align_search_ms=60.0,
):
    pre_s = float(pre_ms) / 1000.0
    post_s = float(post_ms) / 1000.0
    search_s = float(align_search_ms) / 1000.0
    window_start = float(spike_time) - pre_s - search_s
    window_end = float(spike_time) + post_s + search_s
    time_s, signal, fs = _read_spike_window(eeg_module, path, window_start, window_end, fs_hint)
    if time_s is None:
        return None

    detector_idx = int(np.searchsorted(time_s, float(spike_time)))
    detector_idx = max(0, min(detector_idx, len(time_s) - 1))
    align_idx, align_amplitude, align_polarity = _find_alignment_extremum(
        signal,
        detector_idx,
        fs,
        search_ms=align_search_ms,
    )
    align_time = float(time_s[align_idx])

    if target_x_ms is None:
        pre_samples = int(round(pre_ms / 1000.0 * fs))
        post_samples = int(round(post_ms / 1000.0 * fs))
        target_x_ms = (np.arange(-pre_samples, post_samples + 1) / fs) * 1000.0
    else:
        target_x_ms = np.asarray(target_x_ms, dtype=float)

    local_x_ms = (time_s - align_time) * 1000.0
    if local_x_ms[0] > target_x_ms[0] or local_x_ms[-1] < target_x_ms[-1]:
        return None

    segment = np.interp(target_x_ms, local_x_ms, signal).astype(float)
    baseline_points = max(
        3,
        int(np.sum(target_x_ms <= target_x_ms[0] + min(20.0, float(pre_ms)))),
    )
    baseline = float(np.median(segment[:baseline_points]))
    segment = segment - baseline

    return {
        "x_ms": target_x_ms,
        "y": segment,
        "detector_time_s": float(spike_time),
        "align_time_s": align_time,
        "alignment_amplitude": float(align_amplitude),
        "alignment_polarity": align_polarity,
        "fs": float(fs),
    }


def build_average_spike_from_tsvs(file_paths, pre_ms=100.0, post_ms=200.0, progress_callback=None):
    """Detect, align, overlay, and average spikes from multiple TSV/TXT/CSV files."""
    aligned_spikes = []
    event_rows = []
    file_rows = []
    average_x = None
    eeg_module = load_eeg_ecg_analyser_module()

    total_files = len(file_paths)
    for file_index, file_path in enumerate(file_paths, start=1):
        path = Path(file_path)
        declared_samples = _read_declared_sample_count(path)
        last_eeg_progress = {"message": ""}
        if progress_callback:
            progress_callback(
                f"EEG/ECG detection first: file {file_index}/{total_files} {path.name}",
            )

        def eeg_progress(percent, message):
            adjusted_percent = _progress_percent_from_eeg_message(message, declared_samples)
            percent_value = adjusted_percent if adjusted_percent is not None else float(percent)
            formatted = f"{path.name}: {message} ({percent_value:.1f}%)"
            last_eeg_progress["message"] = formatted
            if progress_callback:
                progress_callback(formatted)

        try:
            _, events_df, meta = eeg_module.analyze_eeg_file(
                str(path),
                progress_callback=eeg_progress,
                streaming_threshold_mb=128.0,
            )
        except Exception as exc:
            detail = last_eeg_progress["message"] or "no progress reported"
            raise RuntimeError(
                f"{path.name}: EEG/ECG spike detection failed after {detail}. {exc}"
            ) from exc

        spike_times = _active_spike_times_from_eeg_events(eeg_module, events_df)
        total_detected = int(len(spike_times))
        spike_times = np.sort(spike_times)
        used_count = 0
        skipped_windows = 0
        fs_hint = meta.get("fs_Hz")
        try:
            fs_hint = float(fs_hint)
        except (TypeError, ValueError):
            fs_hint = None

        if progress_callback:
            progress_callback(
                f"{path.name}: EEG/ECG detector found {total_detected} active spikes. Extracting aligned waveforms.",
            )

        progress_step = max(1, total_detected // 20) if total_detected else 1
        for spike_index, spike_time in enumerate(spike_times, start=1):
            if progress_callback and (spike_index == 1 or spike_index % progress_step == 0):
                progress_callback(
                    "Extracting spike windows "
                    f"file {file_index}/{total_files}, spike {spike_index}/{total_detected}, "
                    f"used {used_count} from this file, {len(aligned_spikes)} total",
                )

            try:
                extracted = _extract_aligned_spike_segment(
                    eeg_module,
                    path,
                    float(spike_time),
                    pre_ms,
                    post_ms,
                    target_x_ms=average_x,
                    fs_hint=fs_hint,
                )
            except Exception:
                extracted = None

            if extracted is None:
                skipped_windows += 1
                continue

            if average_x is None:
                average_x = extracted["x_ms"].copy()

            segment = extracted["y"]
            align_time = extracted["align_time_s"]
            align_amplitude = extracted["alignment_amplitude"]
            align_polarity = extracted["alignment_polarity"]

            aligned_spikes.append(
                {
                    "file": str(path),
                    "spike_number": len(aligned_spikes) + 1,
                    "event": {"detector_time_s": float(spike_time), "align_time_s": align_time},
                    "x_ms": average_x.copy(),
                    "y": segment,
                    "raw_y": segment.copy(),
                    "alignment_amplitude": align_amplitude,
                    "alignment_polarity": align_polarity,
                    "polarity_flipped": False,
                }
            )
            event_rows.append(
                {
                    "File": str(path),
                    "Spike": len(aligned_spikes),
                    "Detector time (s)": float(spike_time),
                    "Aligned peak/trough time (s)": align_time,
                    "Alignment shift (ms)": (align_time - float(spike_time)) * 1000.0,
                    "Alignment amplitude": align_amplitude,
                    "Alignment polarity": align_polarity,
                    "Detector": "EEG/ECG analyser",
                    "Saved corrections applied": int(meta.get("eeg_saved_corrections_applied", 0) or 0),
                }
            )
            used_count += 1

        file_rows.append(
            {
                "File": str(path),
                "Rows skipped": 0,
                "Detected spikes": int(total_detected),
                "Spikes used": int(used_count),
                "Waveform windows skipped": int(skipped_windows),
                "EEG saved corrections applied": int(meta.get("eeg_saved_corrections_applied", 0) or 0),
                "EEG detection profile": str(meta.get("eeg_detection_profile", "")),
                "Streaming": bool(meta.get("streaming", False)),
            }
        )

        if progress_callback:
            progress_callback(
                f"Finished {path.name}: {total_detected} detected, {used_count} used, {skipped_windows} skipped.",
            )

    if not aligned_spikes:
        raise ValueError("No complete spikes were detected in the selected TSV files.")

    positive_count = sum(1 for spike in aligned_spikes if spike["alignment_amplitude"] >= 0)
    negative_count = len(aligned_spikes) - positive_count
    dominant_sign = 1.0 if positive_count >= negative_count else -1.0
    dominant_polarity = "positive peak" if dominant_sign > 0 else "negative trough"

    for spike in aligned_spikes:
        spike_sign = 1.0 if spike["alignment_amplitude"] >= 0 else -1.0
        if spike_sign != dominant_sign:
            spike["y"] = -spike["y"]
            spike["polarity_flipped"] = True

    for row in event_rows:
        spike_index = int(row["Spike"]) - 1
        if 0 <= spike_index < len(aligned_spikes):
            row["Dominant average polarity"] = dominant_polarity
            row["Polarity flipped for average"] = aligned_spikes[spike_index]["polarity_flipped"]

    y_stack = np.vstack([spike["y"] for spike in aligned_spikes])
    average_y = np.mean(y_stack, axis=0)

    return {
        "x_ms": average_x,
        "spikes": aligned_spikes,
        "average_y": average_y,
        "event_rows": event_rows,
        "file_rows": file_rows,
        "pre_ms": float(pre_ms),
        "post_ms": float(post_ms),
        "alignment": "local dominant peak/trough within +/-60 ms of EEG/ECG detector event",
        "dominant_polarity": dominant_polarity,
        "positive_alignment_count": int(positive_count),
        "negative_alignment_count": int(negative_count),
    }


def export_average_spikes_to_excel(output_path, average_result):
    """Export every aligned spike plus average spike x,y data."""
    workbook = Workbook()
    all_sheet = workbook.active
    all_sheet.title = "All Spikes"
    all_sheet.append(
        [
            "Spike",
            "Source file",
            "x_ms",
            "y",
            "original_y",
            "Alignment polarity",
            "Polarity flipped for average",
        ]
    )

    for spike in average_result["spikes"]:
        for x_value, y_value, raw_y_value in zip(spike["x_ms"], spike["y"], spike["raw_y"]):
            all_sheet.append(
                [
                    spike["spike_number"],
                    spike["file"],
                    float(x_value),
                    float(y_value),
                    float(raw_y_value),
                    spike["alignment_polarity"],
                    bool(spike["polarity_flipped"]),
                ]
            )
        all_sheet.append([])

    average_sheet = workbook.create_sheet("Average Spike")
    average_sheet.append(["x_ms", "average_y"])
    for x_value, y_value in zip(average_result["x_ms"], average_result["average_y"]):
        average_sheet.append([float(x_value), float(y_value)])

    events_sheet = workbook.create_sheet("Detected Spikes")
    if average_result["event_rows"]:
        headers = list(average_result["event_rows"][0].keys())
        events_sheet.append(headers)
        for row in average_result["event_rows"]:
            events_sheet.append([row.get(header, "") for header in headers])
    else:
        events_sheet.append(["No detected spikes exported"])

    files_sheet = workbook.create_sheet("Files")
    if average_result["file_rows"]:
        headers = list(average_result["file_rows"][0].keys())
    else:
        headers = [
            "File",
            "Rows skipped",
            "Detected spikes",
            "Spikes used",
            "Waveform windows skipped",
            "EEG saved corrections applied",
            "EEG detection profile",
            "Streaming",
        ]
    files_sheet.append(headers)
    for row in average_result["file_rows"]:
        files_sheet.append([row.get(header, "") for header in headers])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.freeze_panes = "A2"
        autosize_worksheet_columns(worksheet)

    workbook.save(output_path)
    return output_path


def save_average_spike_plot(output_path, average_result):
    """Save overlay plot of all detected spikes and their average."""
    figure, axis = plt.subplots(figsize=(10, 6), dpi=150)
    for spike in average_result["spikes"]:
        axis.plot(spike["x_ms"], spike["y"], color="#1f77b4", alpha=0.18, linewidth=1)

    axis.plot(
        average_result["x_ms"],
        average_result["average_y"],
        color="#d7191c",
        linewidth=3,
        label="Average spike",
    )
    axis.axvline(0, color="#111111", linestyle=":", linewidth=1, label="Aligned peak/trough")
    axis.set_title("Peak/trough-aligned epileptiform spikes overlay")
    axis.set_xlabel("Time from aligned peak/trough (ms)")
    axis.set_ylabel("Baseline-corrected signal (polarity-normalized)")
    axis.grid(True)
    axis.legend()
    figure.tight_layout()
    figure.savefig(output_path)
    plt.close(figure)
    return output_path


def load_eeg_ecg_analyser_module():
    """Load the bundled EEG/ECG helper used by average-spike analysis."""
    try:
        import eeg_ecg_analyser_shared as module
    except Exception as exc:
        raise RuntimeError("Bundled EEG/ECG analyser helper could not be loaded.") from exc
    return module


def _active_spike_times_from_eeg_events(eeg_module, events_df):
    if hasattr(eeg_module, "active_eeg_events"):
        active = eeg_module.active_eeg_events(events_df)
    else:
        active = events_df.copy()
        if "Is_Deleted" in active.columns:
            active = active.loc[~active["Is_Deleted"].astype(bool)]

    if len(active) == 0 or "Type" not in active.columns:
        return np.array([], dtype=float)

    spike_rows = active.loc[active["Type"].astype(str) == "Spike"]
    if "Start_s" not in spike_rows.columns:
        return np.array([], dtype=float)

    return spike_rows["Start_s"].to_numpy(dtype=float)


def format_number(value, digits=10):
    """Format numbers for readable equations without unnecessary trailing zeros."""
    text = f"{value:.{digits}g}"
    if text == "-0":
        return "0"
    return text


def format_signed(value, digits=10):
    """Format a signed number for equations."""
    sign = "+" if value >= 0 else "-"
    return f"{sign} {format_number(abs(value), digits)}"


def format_sinusoid_equation(A, B, C, D, digits=6):
    """Format y = A*sin(B*x + C) + D as readable text."""
    c_sign = "+" if C >= 0 else "-"
    d_sign = "+" if D >= 0 else "-"
    return (
        f"y = {format_number(A, digits)} * sin("
        f"{format_number(B, digits)} * x {c_sign} {format_number(abs(C), digits)}) "
        f"{d_sign} {format_number(abs(D), digits)}"
    )


def format_spline_formula(a, b, c, d, x0, digits=10):
    """Format one cubic spline interval as readable text."""
    return (
        f"y = {format_number(a, digits)}*(x - {format_number(x0, digits)})^3 "
        f"{format_signed(b, digits)}*(x - {format_number(x0, digits)})^2 "
        f"{format_signed(c, digits)}*(x - {format_number(x0, digits)}) "
        f"{format_signed(d, digits)}"
    )


def build_spline_formula_rows(x_values, spline):
    """Build formula rows for Excel export."""
    rows = []

    for index in range(len(x_values) - 1):
        x0 = float(x_values[index])
        x1 = float(x_values[index + 1])
        a, b, c, d = [float(value) for value in spline.c[:, index]]

        rows.append(
            {
                "Interval": index + 1,
                "x_start": x0,
                "x_end": x1,
                "a": a,
                "b": b,
                "c": c,
                "d": d,
                "Formula": format_spline_formula(a, b, c, d, x0),
            }
        )

    return rows


def build_sinusoid_formula_rows(A, B, C, D):
    """Build formula rows for sinusoidal Excel export."""
    return [
        {
            "Formula Type": "Sinusoidal best fit",
            "A": float(A),
            "B": float(B),
            "C": float(C),
            "D": float(D),
            "Formula": format_sinusoid_equation(A, B, C, D),
        }
    ]


def autosize_worksheet_columns(worksheet):
    """Adjust Excel column widths to fit exported content."""
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 90)


def export_analysis_to_excel(
    output_path,
    metadata_rows,
    area_rows,
    formula_rows,
    spike_metric_rows=None,
    data_rows=None,
):
    """Export area results, formulas, spike metrics, and optional x/y data."""
    workbook = Workbook()

    area_sheet = workbook.active
    area_sheet.title = "Area Results"
    area_sheet.append(["Field", "Value"])

    for key, value in metadata_rows:
        area_sheet.append([key, value])

    area_sheet.append([])
    area_sheet.append(["Area Result", "Value"])

    for key, value in area_rows:
        area_sheet.append([key, value])

    formulas_sheet = workbook.create_sheet("Formulas")

    if formula_rows:
        headers = list(formula_rows[0].keys())
        formulas_sheet.append(headers)

        for row in formula_rows:
            formulas_sheet.append([row.get(header, "") for header in headers])
    else:
        formulas_sheet.append(["Formula"])
        formulas_sheet.append(["No formula rows were generated."])

    if spike_metric_rows:
        spike_sheet = workbook.create_sheet("Spike Metrics")
        headers = list(spike_metric_rows[0].keys())
        spike_sheet.append(headers)

        for row in spike_metric_rows:
            spike_sheet.append([row.get(header, "") for header in headers])

    if data_rows:
        data_sheet = workbook.create_sheet("Data")
        headers = list(data_rows[0].keys())
        data_sheet.append(headers)

        for row in data_rows:
            data_sheet.append([row.get(header, "") for header in headers])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)

        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        worksheet.freeze_panes = "A2"
        autosize_worksheet_columns(worksheet)

    workbook.save(output_path)
    return output_path


def build_data_rows(x_values, y_values):
    """Build x/y data rows for analysis workbook reloads."""
    return [
        {
            "x": float(x_value),
            "y": float(y_value),
        }
        for x_value, y_value in zip(x_values, y_values)
    ]


def _sheet_to_key_values(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        return []

    rows = []
    worksheet = workbook[sheet_name]
    in_area = False
    for row in worksheet.iter_rows(values_only=True):
        if not row or row[0] in (None, ""):
            continue
        if row[0] == "Area Result":
            in_area = True
            continue
        if in_area:
            continue
        key = str(row[0])
        if key in {"Field", "Value", "Area Result", "Metric"}:
            continue
        rows.append((key, row[1] if len(row) > 1 else ""))

    return rows


def _read_table_rows(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        return []

    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = next(iterator)
    except StopIteration:
        return []

    headers = [str(header) if header is not None else "" for header in headers]
    rows = []
    for values in iterator:
        if not values or all(value in (None, "") for value in values):
            continue
        rows.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
        )

    return rows


def _rows_to_xy(rows, x_key, y_key):
    x_values = []
    y_values = []
    for row in rows:
        x_value = to_float(row.get(x_key))
        y_value = to_float(row.get(y_key))
        if x_value is None or y_value is None:
            continue
        x_values.append(x_value)
        y_values.append(y_value)

    if len(x_values) < 4:
        return None, None

    return np.array(x_values, dtype=float), np.array(y_values, dtype=float)


def load_previous_analysis_workbook(file_path):
    """Load a previous Curve Analyzer workbook for viewing."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        metadata_rows = _sheet_to_key_values(workbook, "Area Results")
        area_rows = []
        if "Area Results" in workbook.sheetnames:
            area_sheet = workbook["Area Results"]
            in_area = False
            for row in area_sheet.iter_rows(values_only=True):
                if not row or row[0] in (None, ""):
                    continue
                if row[0] == "Area Result":
                    in_area = True
                    continue
                if in_area and row[0] != "Value":
                    area_rows.append((str(row[0]), row[1] if len(row) > 1 else ""))

        spike_metric_rows = _read_table_rows(workbook, "Spike Metrics")
        formula_rows = _read_table_rows(workbook, "Formulas")
        source_file = Path(file_path)

        if "Average Spike" in workbook.sheetnames:
            average_rows = _read_table_rows(workbook, "Average Spike")
            x_values, y_values = _rows_to_xy(average_rows, "x_ms", "average_y")
            if x_values is None:
                raise ValueError("Average Spike sheet does not contain usable x_ms/average_y data.")

            overlay_rows = _read_table_rows(workbook, "All Spikes")
            overlay_spikes = []
            grouped = {}
            for row in overlay_rows:
                spike_number = row.get("Spike")
                x_value = to_float(row.get("x_ms"))
                y_value = to_float(row.get("y"))
                if spike_number is None or x_value is None or y_value is None:
                    continue
                grouped.setdefault(spike_number, {"x": [], "y": []})
                grouped[spike_number]["x"].append(x_value)
                grouped[spike_number]["y"].append(y_value)

            for spike_number, values in grouped.items():
                if len(values["x"]) >= 4:
                    overlay_spikes.append(
                        {
                            "spike_number": spike_number,
                            "x_ms": np.array(values["x"], dtype=float),
                            "y": np.array(values["y"], dtype=float),
                        }
                    )

            x_values, y_values, spline = make_exact_spline(x_values, y_values)
            signed_area, absolute_area = calculate_spline_area(spline, x_values[0], x_values[-1])
            spike_metrics = calculate_spike_metrics(x_values, y_values)
            return {
                "mode": "previous_average_spike",
                "source_file": source_file,
                "sheet_name": "Average Spike",
                "x": x_values,
                "y": y_values,
                "spline": spline,
                "signed_area": signed_area,
                "absolute_area": absolute_area,
                "metadata_rows": metadata_rows,
                "area_rows": area_rows or [
                    ("x start", float(x_values[0])),
                    ("x end", float(x_values[-1])),
                    ("Signed area under curve", signed_area),
                    ("Absolute area under curve", absolute_area),
                ],
                "formula_rows": formula_rows,
                "spike_metrics": spike_metrics,
                "spike_metric_rows": spike_metric_rows or build_spike_metric_rows(spike_metrics),
                "overlay_spikes": overlay_spikes,
                "data_rows": build_data_rows(x_values, y_values),
            }

        if "Data" in workbook.sheetnames:
            data_rows = _read_table_rows(workbook, "Data")
            x_values, y_values = _rows_to_xy(data_rows, "x", "y")
            if x_values is None:
                raise ValueError("Data sheet does not contain usable x/y data.")

            x_values, y_values, spline = make_exact_spline(x_values, y_values)
            signed_area, absolute_area = calculate_spline_area(spline, x_values[0], x_values[-1])
            spike_metrics = calculate_spike_metrics(x_values, y_values)
            return {
                "mode": "previous_data",
                "source_file": source_file,
                "sheet_name": "Data",
                "x": x_values,
                "y": y_values,
                "spline": spline,
                "signed_area": signed_area,
                "absolute_area": absolute_area,
                "metadata_rows": metadata_rows,
                "area_rows": area_rows or [
                    ("x start", float(x_values[0])),
                    ("x end", float(x_values[-1])),
                    ("Signed area under curve", signed_area),
                    ("Absolute area under curve", absolute_area),
                ],
                "formula_rows": formula_rows,
                "spike_metrics": spike_metrics,
                "spike_metric_rows": spike_metric_rows or build_spike_metric_rows(spike_metrics),
                "data_rows": build_data_rows(x_values, y_values),
            }

        if metadata_rows or area_rows or spike_metric_rows or formula_rows:
            return {
                "mode": "previous_summary",
                "source_file": source_file,
                "sheet_name": "Workbook Summary",
                "metadata_rows": metadata_rows,
                "area_rows": area_rows,
                "formula_rows": formula_rows,
                "spike_metric_rows": spike_metric_rows,
            }

        raise ValueError("This workbook does not look like a Curve Analyzer result workbook.")
    finally:
        workbook.close()


def save_spline_equations(
    x_values,
    spline,
    output_path,
    signed_area=None,
    absolute_area=None,
):
    """Save the exact piecewise cubic spline equations to a text file."""
    with open(output_path, "w", encoding="utf-8") as file:
        file.write("Exact interpolated cubic spline formula\n")
        file.write("Column A is x. Column B is y.\n\n")

        if signed_area is not None:
            file.write(
                f"Signed area from x = {format_number(x_values[0])} "
                f"to x = {format_number(x_values[-1])}: "
                f"{format_number(signed_area)}\n"
            )

        if absolute_area is not None:
            file.write(
                f"Absolute area from x = {format_number(x_values[0])} "
                f"to x = {format_number(x_values[-1])}: "
                f"{format_number(absolute_area)}\n"
            )

        if signed_area is not None or absolute_area is not None:
            file.write("\n")

        file.write(
            "Each interval uses this form:\n"
            "y = a*(x - x0)^3 + b*(x - x0)^2 + c*(x - x0) + d\n\n"
        )

        for index in range(len(x_values) - 1):
            x0 = x_values[index]
            x1 = x_values[index + 1]
            a, b, c, d = spline.c[:, index]

            file.write(
                f"For {format_number(x0)} <= x <= {format_number(x1)}:\n"
            )
            file.write(f"{format_spline_formula(a, b, c, d, x0)}\n\n")


def print_spline_summary(x_values, output_path, signed_area, absolute_area):
    """Print a readable summary for the exact interpolated formula."""
    print("Exact interpolated formula:")
    print("A cubic spline was created through every data point.")
    print(f"Number of piecewise equations: {len(x_values) - 1}")
    print(f"Formula file saved to: {output_path}")
    print()
    print(f"Area range: x = {x_values[0]:.6f} to x = {x_values[-1]:.6f}")
    print(f"Signed area under curve:   {signed_area:.6f}")
    print(f"Absolute area under curve: {absolute_area:.6f}")
    print()
    print("Important:")
    print("This is exact at the Excel points, but it is a piecewise formula.")
    print("There is one cubic equation between each pair of x values.")


def print_spike_summary(spike_metrics):
    """Print dominant spike morphology metrics."""
    print()
    print("Dominant spike metrics:")
    print(f"Baseline y: {spike_metrics['Baseline y (median)']:.6f}")
    print(f"Polarity: {spike_metrics['Spike polarity']}")
    print(f"Peak x: {spike_metrics['Peak x']:.6f}")
    print(f"Peak y: {spike_metrics['Peak y']:.6f}")
    print(
        "Peak amplitude from baseline: "
        f"{spike_metrics['Peak amplitude (absolute from baseline)']:.6f}"
    )

    rise_time = spike_metrics["Rise time (10-90%)"]
    decay_time = spike_metrics["Decay time (90-10%)"]
    print(
        "Rise time (10-90%): "
        f"{rise_time:.6f}" if rise_time is not None else "Rise time (10-90%): not found"
    )
    print(
        "Decay time (90-10%): "
        f"{decay_time:.6f}" if decay_time is not None else "Decay time (90-10%): not found"
    )


def to_float(value):
    """Convert worksheet values to floats, returning None for headers/blanks."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_xy_from_excel(file_path):
    """Load x values from column A and y values from column B."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active
    worksheet_title = worksheet.title

    points = []
    skipped_rows = 0

    for x_raw, y_raw in worksheet.iter_rows(
        min_col=1,
        max_col=2,
        values_only=True,
    ):
        x_value = to_float(x_raw)
        y_value = to_float(y_raw)

        if x_value is None or y_value is None:
            if x_raw is not None or y_raw is not None:
                skipped_rows += 1
            continue

        points.append((x_value, y_value))

    workbook.close()

    if len(points) < 4:
        raise ValueError(
            "The Excel sheet needs at least 4 numeric rows in columns A and B."
        )

    points.sort(key=lambda point: point[0])

    x_values = np.array([point[0] for point in points], dtype=float)
    y_values = np.array([point[1] for point in points], dtype=float)

    return x_values, y_values, worksheet_title, len(points), skipped_rows


def choose_excel_file():
    """Open a file picker so the user can select an Excel workbook."""
    try:
        from tkinter import Tk, filedialog
    except ImportError:
        return None

    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    file_path = filedialog.askopenfilename(
        title="Select Excel file: column A = x, column B = y",
        filetypes=[
            ("Excel workbooks", "*.xlsx *.xlsm"),
            ("All files", "*.*"),
        ],
    )

    root.destroy()
    return file_path or None


def choose_analysis_mode(default_mode="exact"):
    """Ask whether to use sinusoidal best fit or exact interpolation."""
    print("Choose graph formula mode:")
    print("1 = Best-fit sinusoidal equation")
    print("2 = Exact interpolated curve through every point")

    default_choice = "2" if default_mode == "exact" else "1"
    choice = input(f"Enter 1 or 2 [{default_choice}]: ").strip()

    if not choice:
        choice = default_choice

    if choice == "1":
        return "sinusoid"

    if choice == "2":
        return "exact"

    print("Invalid choice. Using exact interpolated curve.")
    return "exact"


def parse_arguments():
    """Parse optional command-line arguments."""
    excel_path = None
    mode = None

    for argument in sys.argv[1:]:
        normalized = argument.lower().strip()

        if normalized in {"--exact", "--interpolate", "--spline"}:
            mode = "exact"
        elif normalized in {"--sinusoid", "--best-fit", "--bestfit"}:
            mode = "sinusoid"
        else:
            excel_path = Path(argument)

    return excel_path, mode


def print_equation(
    A,
    B,
    C,
    D,
    signed_area=None,
    absolute_area=None,
    x_min=None,
    x_max=None,
):
    """Print the fitted sinusoidal equation in readable form."""
    print("Best-fit parameters:")
    print(f"Amplitude A:       {A:.6f}")
    print(f"Frequency B:       {B:.6f}")
    print(f"Phase shift C:     {C:.6f}")
    print(f"Vertical offset D: {D:.6f}")
    print()
    print("Best-fit sinusoidal equation:")
    print(format_sinusoid_equation(A, B, C, D, digits=4))

    if signed_area is not None:
        print()
        print(f"Area range: x = {x_min:.6f} to x = {x_max:.6f}")
        print(f"Signed area under fitted curve: {signed_area:.6f}")
        if absolute_area is not None:
            print(f"Absolute area under fitted curve: {absolute_area:.6f}")


def plot_fit(
    x_values,
    y_values,
    A,
    B,
    C,
    D,
    title="Sinusoidal Curve Fit",
    signed_area=None,
):
    """Plot the original data points and fitted sinusoidal curve."""
    x_fit = np.linspace(np.min(x_values), np.max(x_values), 1000)
    y_fit = sinusoid(x_fit, A, B, C, D)

    c_sign = "+" if C >= 0 else "-"
    d_sign = "+" if D >= 0 else "-"
    equation_text = (
        f"y = {A:.4f} sin({B:.4f}x {c_sign} {abs(C):.4f}) "
        f"{d_sign} {abs(D):.4f}"
    )

    if signed_area is not None:
        equation_text += f"\nSigned area = {signed_area:.4f}"

    plt.scatter(x_values, y_values, label="Original data", color="blue")
    plt.plot(x_fit, y_fit, label="Fitted curve", color="red", linewidth=2)
    plt.fill_between(
        x_fit,
        y_fit,
        0,
        color="red",
        alpha=0.12,
        label="Area under fitted curve",
    )

    plt.xlabel("x")
    plt.ylabel("y")
    plt.title(title)
    note = plt.text(
        0.02,
        0.98,
        equation_text,
        transform=plt.gca().transAxes,
        verticalalignment="top",
        bbox={"facecolor": "white", "alpha": 0.8, "edgecolor": "gray"},
    )
    note.set_picker(True)
    legend = plt.legend()
    legend.set_draggable(True)
    plt.grid(True)
    plt.show()


def plot_exact_spline(
    x_values,
    y_values,
    spline,
    title="Exact Interpolated Curve",
    signed_area=None,
    absolute_area=None,
):
    """Plot original points and the exact interpolated spline curve."""
    x_fit = np.linspace(np.min(x_values), np.max(x_values), 2000)
    y_fit = spline(x_fit)

    label_text = "Exact cubic spline interpolation\nFormula saved as piecewise equations"
    if signed_area is not None:
        label_text += f"\nSigned area = {signed_area:.4f}"
    if absolute_area is not None:
        label_text += f"\nAbsolute area = {absolute_area:.4f}"

    plt.scatter(x_values, y_values, label="Original data", color="blue")
    plt.plot(
        x_fit,
        y_fit,
        label="Exact interpolated curve",
        color="red",
        linewidth=2,
    )
    plt.fill_between(
        x_fit,
        y_fit,
        0,
        color="red",
        alpha=0.12,
        label="Area under curve",
    )

    plt.xlabel("x")
    plt.ylabel("y")
    plt.title(title)
    note = plt.text(
        0.02,
        0.98,
        label_text,
        transform=plt.gca().transAxes,
        verticalalignment="top",
        bbox={"facecolor": "white", "alpha": 0.8, "edgecolor": "gray"},
    )
    note.set_picker(True)
    legend = plt.legend()
    legend.set_draggable(True)
    plt.grid(True)
    plt.show()


def make_sample_data():
    """Create sample data if no Excel file is selected."""
    np.random.seed(42)

    x_data = np.linspace(0, 10, 80)

    true_A = 2.5
    true_B = 1.7
    true_C = 0.6
    true_D = 1.2

    noise = np.random.normal(0, 0.3, size=len(x_data))
    y_data = sinusoid(x_data, true_A, true_B, true_C, true_D) + noise

    return x_data, y_data


def main():
    excel_path, mode = parse_arguments()

    if excel_path is None:
        selected_file = choose_excel_file()
        excel_path = Path(selected_file) if selected_file else None

    if excel_path:
        if not excel_path.exists():
            raise FileNotFoundError(f"Could not find Excel file: {excel_path}")

        x_data, y_data, sheet_name, used_rows, skipped_rows = load_xy_from_excel(
            excel_path
        )

        print(f"Loaded file: {excel_path}")
        print(f"Worksheet: {sheet_name}")
        print(f"Numeric data rows used: {used_rows}")
        print(f"Rows skipped: {skipped_rows}")
        print()

        base_title = excel_path.name
        output_stem = excel_path.with_suffix("")
        source_label = str(excel_path)
    else:
        print("No Excel file selected. Using built-in sample data.")
        print()
        x_data, y_data = make_sample_data()
        sheet_name = "Built-in sample"
        used_rows = len(x_data)
        skipped_rows = 0
        base_title = "Sample Data"
        output_stem = Path("sample_data")
        source_label = "Built-in sample data"

    if mode is None:
        mode = choose_analysis_mode(default_mode="exact")

    spike_metrics = calculate_spike_metrics(x_data, y_data)
    spike_metric_rows = build_spike_metric_rows(spike_metrics)

    if mode == "sinusoid":
        A, B, C, D = fit_sinusoidal(x_data, y_data)
        x_min = float(np.min(x_data))
        x_max = float(np.max(x_data))
        signed_area = calculate_sinusoid_area(A, B, C, D, x_min, x_max)
        absolute_area = calculate_sinusoid_absolute_area(A, B, C, D, x_min, x_max)
        results_path = output_stem.with_name(f"{output_stem.name}_analysis_results.xlsx")

        print_equation(
            A,
            B,
            C,
            D,
            signed_area=signed_area,
            absolute_area=absolute_area,
            x_min=x_min,
            x_max=x_max,
        )

        export_analysis_to_excel(
            results_path,
            metadata_rows=[
                ("Source file", source_label),
                ("Worksheet", sheet_name),
                ("Analysis mode", "Sinusoidal best fit"),
                ("Numeric data rows used", used_rows),
                ("Rows skipped", skipped_rows),
                ("Formula type", "y = A * sin(B*x + C) + D"),
            ],
            area_rows=[
                ("x start", x_min),
                ("x end", x_max),
                ("Signed area under curve", signed_area),
                ("Absolute area under curve", absolute_area),
            ],
            formula_rows=build_sinusoid_formula_rows(A, B, C, D),
            spike_metric_rows=spike_metric_rows,
            data_rows=build_data_rows(x_data, y_data),
        )
        print_spike_summary(spike_metrics)
        print(f"Excel results saved to: {results_path}")

        plot_fit(
            x_data,
            y_data,
            A,
            B,
            C,
            D,
            title=f"Sinusoidal Best Fit: {base_title}",
            signed_area=signed_area,
        )
    else:
        x_data, y_data, spline = make_exact_spline(x_data, y_data)
        formula_path = output_stem.with_name(f"{output_stem.name}_exact_formula.txt")
        results_path = output_stem.with_name(f"{output_stem.name}_analysis_results.xlsx")
        signed_area, absolute_area = calculate_spline_area(
            spline,
            x_data[0],
            x_data[-1],
        )

        save_spline_equations(
            x_data,
            spline,
            formula_path,
            signed_area=signed_area,
            absolute_area=absolute_area,
        )
        export_analysis_to_excel(
            results_path,
            metadata_rows=[
                ("Source file", source_label),
                ("Worksheet", sheet_name),
                ("Analysis mode", "Exact cubic spline interpolation"),
                ("Numeric data rows used", used_rows),
                ("Rows skipped", skipped_rows),
                ("Unique x values", len(x_data)),
                (
                    "Formula type",
                    "Piecewise cubic spline: y = a*(x-x0)^3 + b*(x-x0)^2 + c*(x-x0) + d",
                ),
            ],
            area_rows=[
                ("x start", float(x_data[0])),
                ("x end", float(x_data[-1])),
                ("Signed area under curve", signed_area),
                ("Absolute area under curve", absolute_area),
            ],
            formula_rows=build_spline_formula_rows(x_data, spline),
            spike_metric_rows=spike_metric_rows,
            data_rows=build_data_rows(x_data, y_data),
        )
        print_spline_summary(x_data, formula_path, signed_area, absolute_area)
        print_spike_summary(spike_metrics)
        print(f"Excel results saved to: {results_path}")
        plot_exact_spline(
            x_data,
            y_data,
            spline,
            title=f"Exact Interpolated Curve: {base_title}",
            signed_area=signed_area,
            absolute_area=absolute_area,
        )


if __name__ == "__main__":
    main()
